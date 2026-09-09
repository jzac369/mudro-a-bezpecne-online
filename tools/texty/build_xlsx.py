# -*- coding: utf-8 -*-
"""Zostaví pracovný hárok s textami kurzu: pôvodné znenie, návrh a miesto na finálnu úpravu."""
import json, sys, re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

texts_path, proposals_path, out_path = sys.argv[1], sys.argv[2], sys.argv[3]
rows = json.load(open(texts_path, encoding="utf-8"))
proposals = json.load(open(proposals_path, encoding="utf-8")) if proposals_path != "-" else {}

FONT = "Arial"
INK = "1F3A3D"
wb = Workbook()

# ---------------------------------------------------------------- Návod
gu = wb.active
gu.title = "Návod"
gu.sheet_view.showGridLines = False
guide = [
    ("Texty lekcií — pracovný hárok", 16, True, INK),
    ("", 11, False, None),
    ("Na čo to je", 12, True, INK),
    ("Hárok „Texty kurzu“ obsahuje KAŽDÝ text, ktorý účastník v Lekciách uvidí — nadpisy, zadania,", 11, False, None),
    ("vysvetlenia pri odpovediach, poučky aj popisky. Sú to presne tie polia, ktoré sa dajú upraviť", 11, False, None),
    ("aj v admin zóne, takže čokoľvek odtiaľto viem vrátiť späť do systému.", 11, False, None),
    ("", 11, False, None),
    ("Ako s tým pracovať", 12, True, INK),
    ("1. Prečítajte si stĺpec D „Pôvodný text“ a vedľa v stĺpci E „Návrh“ moju zrozumiteľnejšiu verziu.", 11, False, None),
    ("2. Ak vám návrh vyhovuje, nechajte stĺpec F prázdny — použije sa návrh.", 11, False, None),
    ("3. Ak chcete niečo doladiť, napíšte svoje znenie do stĺpca F „Finálny text“. Ten má vždy prednosť.", 11, False, None),
    ("4. Ak má text zostať tak, ako je, napíšte do stĺpca G „Poznámka“ slovo PONECHAŤ.", 11, False, None),
    ("", 11, False, None),
    ("Upravujte iba stĺpce F a G. Stĺpce A až E slúžia na priradenie textu späť do systému —", 11, True, "9C4A1F"),
    ("keď sa zmenia, nebudem vedieť, kam text patrí.", 11, True, "9C4A1F"),
    ("", 11, False, None),
    ("Príklad vyplnenia", 12, True, INK),
    ("Pôvodný text:  Overte si to.", 11, False, None),
    ("Návrh:         Skôr než niečo urobíte, overte si informáciu na oficiálnej stránke alebo", 11, False, None),
    ("               telefónnom čísle, ktoré si nájdete sami — nie na tom z e-mailu.", 11, False, None),
    ("Finálny text:  Skôr než niečo urobíte, zavolajte do banky na číslo z vašej karty.", 11, False, None),
    ("Poznámka:      (prázdne)", 11, False, None),
    ("", 11, False, None),
    ("Keď budete hotový, pošlite mi súbor späť — texty doplním do kurzu a nasadím.", 11, True, INK),
]
for i, (txt, size, bold, color) in enumerate(guide, start=1):
    c = gu.cell(row=i, column=1, value=txt)
    c.font = Font(name=FONT, size=size, bold=bold, color=color or "000000")
gu.column_dimensions["A"].width = 105

# ---------------------------------------------------------------- Texty
ws = wb.create_sheet("Texty kurzu")
headers = ["Obrazovka", "Pole", "Cesta v systéme", "Pôvodný text", "Návrh (Claude)", "Finálny text (vyplňte)", "Poznámka"]
head_fill = PatternFill("solid", fgColor="12554A")
for i, h in enumerate(headers, start=1):
    c = ws.cell(row=1, column=i, value=h)
    c.font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
    c.fill = head_fill
    c.alignment = Alignment(vertical="center", wrap_text=True)
ws.row_dimensions[1].height = 30

widths = [26, 20, 30, 58, 58, 58, 18]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

thin = Side(style="thin", color="DDD5C2")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
band = PatternFill("solid", fgColor="F3F0E8")
edit_fill = PatternFill("solid", fgColor="FFFCE8")
todo_font = Font(name=FONT, size=10, italic=True, color="8A7F66")

seen_screens = []
r = 2
for row in rows:
    if row["slideId"] not in seen_screens:
        seen_screens.append(row["slideId"])
    shade = band if (seen_screens.index(row["slideId"]) % 2 == 1) else None
    key = str(row["slideId"]) + "|" + row["path"] + "|" + row["text"]
    proposal = proposals.get(key, "")

    values = [
        "%d · %s" % (row["slideId"], row["title"]),
        row["label"],
        row["path"],
        row["text"],
        proposal if proposal else "— doplním —",
        "",
        "",
    ]
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = Font(name=FONT, size=10)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        c.border = border
        if shade and i <= 5:
            c.fill = shade
        if i in (6, 7):
            c.fill = edit_fill
    if not proposal:
        ws.cell(row=r, column=5).font = todo_font
    ws.cell(row=r, column=3).font = Font(name=FONT, size=8, color="8A7F66")
    r += 1

ws.freeze_panes = "D2"
ws.auto_filter.ref = "A1:G%d" % (r - 1)
wb.save(out_path)
print("riadkov:", r - 2, "->", out_path)
